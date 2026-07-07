#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
체육과 교수학습 및 평가 운영 계획 DOCX 생성기
사용법: python3 make_docx.py <입력.json> [출력.docx]

입력 JSON 구조는 claude-project/지식/JSON_스키마.md 참고.
(meta + part1: 평가계획 / part2: 진도표 / part3: 성취수준)
"""
import json
import sys

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor, Twips

FONT = '맑은 고딕'
HEADER_FILL = 'D9E2F3'   # 표 머리글 배경
SECTION_FILL = 'EEF3FB'  # 부(部) 제목 배경
ACCENT = '2E74B5'        # 부 제목 강조색
CW = 16838 - 1134 * 2    # 본문 폭(twips): A3 세로 - 좌우 여백 2cm

방침 = [
    '가. 체육과는 신체활동의 학습을 통해 활동적이고 창의적인 삶, 건강하고 주도적인 삶, 신체활동 문화를 향유하며 더불어 사는 삶의 자질을 기르는 것을 목적으로 하며, 평가는 이러한 교과의 목적에 부합하도록 신체활동 역량의 전반적인 성장을 확인하는 방향으로 실시한다.',
    '나. 지식·이해, 과정·기능, 가치·태도의 세 가지 내용 요소를 균형 있게 평가하되, 신체활동 수행 과정에서 드러나는 학생의 변화와 성장에 중점을 두어 과정 중심 평가를 실시한다.',
    '다. 평가 결과는 학생의 현재 수준을 파악하고 다음 학습 단계로의 성장을 지원하는 피드백 자료로 활용하며, 수업과 평가가 긴밀하게 연계되도록 한다.',
    '라. 학습자의 실생활 맥락과 연계한 수행평가를 설계‧실시하고 수업 과정에서 형성평가를 활용하여 학생의 학습 상태를 파악하며, 그 결과를 바탕으로 개별 학습 수준에 맞는 맞춤형 피드백과 지원을 제공하여 모든 학생이 학습 목표에 도달하도록 한다.',
    '마. 평가 과정에 학생이 주체적으로 참여하도록 하고, 학습 목표와 수행 과정을 명확히 제시하며 현재 수준과 개선 방향에 대한 구체적인 피드백을 제공함으로써 성장 중심 평가를 구현한다.',
    '바. 수행평가의 전 과정에서 학문적 정직성을 기반으로 학생의 성장을 지원하는 평가 환경을 조성한다. 평가 과정에서 타인의 결과물을 그대로 사용하는 행위 및 교사가 안내한 AI의 허용 범위를 벗어난 활용, 학생 간 공모 및 담합, 기존 제출물의 재사용 등 학문적 정직성 위반 행위에 대해 평소 지속적으로 교육하고, 평가 실시 전 이를 재안내한다.',
]

유의사항 = [
    '가. 지필고사 결시생은 본교 학업성적관리규정에 의거하여 학업성적관리위원회 심의를 거쳐 처리한다.',
    '나. 수행평가 결시생인 경우 추후 평가를 실시함을 원칙으로 한다.',
    '다. 재취학, 전·편입학생인 경우 수행평가 실시 가능 여부에 따라 처리한다.',
    '라. 건강장애로 수행평가에 참여하기 어려운 경우 동일 과제나 대체 과제를 제시한다.',
    '마. 위 조항을 적용하기 어려운 경우 교과협의회 및 학업성적관리위원회 심의를 거쳐 결정한다.',
]


def set_font(run, size=9, bold=False, color=None):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.append(rfonts)
    rfonts.set(qn('w:eastAsia'), FONT)


def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:fill'), fill)
    tcpr.append(shd)


def para(doc, text='', size=9, bold=False, align='left', after=5, indent=0):
    p = doc.add_paragraph()
    p.alignment = {'left': WD_ALIGN_PARAGRAPH.LEFT,
                   'center': WD_ALIGN_PARAGRAPH.CENTER}[align]
    p.paragraph_format.space_after = Pt(after)
    if indent:
        p.paragraph_format.left_indent = Twips(indent)
    set_font(p.add_run(str(text)), size=size, bold=bold)
    return p


def fill_cell(cell, text, size=9, bold=False, align='left', fill=None,
              valign='center', color=None):
    """셀에 텍스트(개행 포함 가능)를 채운다."""
    cell.vertical_alignment = (WD_ALIGN_VERTICAL.TOP if valign == 'top'
                               else WD_ALIGN_VERTICAL.CENTER)
    if fill:
        shade(cell, fill)
    lines = str(text).split('\n') if text is not None else ['']
    first = True
    for line in lines:
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        p.alignment = {'left': WD_ALIGN_PARAGRAPH.LEFT,
                       'center': WD_ALIGN_PARAGRAPH.CENTER}[align]
        p.paragraph_format.space_after = Pt(0)
        set_font(p.add_run(line), size=size, bold=bold, color=color)


def header_cell(cell, text, size=9):
    fill_cell(cell, text, size=size, bold=True, align='center', fill=HEADER_FILL)


def make_table(doc, widths, rows):
    """widths: twips 목록. rows: 행 수."""
    t = doc.add_table(rows=rows, cols=len(widths))
    t.style = 'Table Grid'
    t.autofit = False
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Twips(w)
    return t


def section_header(doc, num, title):
    t = make_table(doc, [CW], 1)
    cell = t.rows[0].cells[0]
    shade(cell, SECTION_FILL)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    set_font(p.add_run(f' {num}  '), size=12, bold=True, color=ACCENT)
    set_font(p.add_run(title), size=12, bold=True)
    para(doc, '', after=5)


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def build(data, out_path):
    meta = data.get('meta', {})
    year = meta.get('year', '2026')
    semester = meta.get('semester', '1')
    school = meta.get('school', '')
    grade = meta.get('grade', '')
    teacher = meta.get('teacher', '')
    eval_period = meta.get('evalPeriod', '')
    sport = meta.get('sport', '체육')

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(29.7), Cm(42.0)  # A3 세로
    for m in ('top_margin', 'bottom_margin', 'left_margin', 'right_margin'):
        setattr(sec, m, Cm(2))
    style = doc.styles['Normal']
    style.font.name = FONT
    style.font.size = Pt(9)
    style.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)

    # 제목
    para(doc, f'{year}학년도 {semester}학기 (체육)과 교수학습 및 평가 운영 계획',
         size=13, bold=True, align='center', after=9)

    # 기본 정보
    bw = [1800, 900, 900, 1300, 4470, 5200]
    t = make_table(doc, bw, 2)
    heads = ['학교명', '학년', '시수', '성취도', '기준학급(수업 요일)', '지도교사']
    vals = [school, f'{grade}학년', '', '3단계', '', teacher]
    for i, h in enumerate(heads):
        header_cell(t.rows[0].cells[i], h)
        fill_cell(t.rows[1].cells[i], vals[i], align='center')
    para(doc, '', after=8)

    # ── 1부: 평가 계획 ──────────────────────────────
    section_header(doc, '1', '(체육)과 평가 계획')
    para(doc, '1. 평가 목적 및 평가 방향, 평가 방침', size=10, bold=True, after=4)
    t = make_table(doc, [CW], 1)
    fill_cell(t.rows[0].cells[0], '\n'.join(방침), valign='top', fill='F8FAFF')
    para(doc, '', after=6)

    # 평가 개요
    para(doc, '2. 평가 개요', size=10, bold=True, after=4)
    eps = data.get('part1', {}).get('evalPlans', [])
    n = max(len(eps), 1)
    per = (CW - 1700) // n
    ew = [1700] + [per] * (n - 1) + [CW - 1700 - per * (n - 1)]
    t = make_table(doc, ew, 5)
    labels = ['평가 유형', '반영 비율', '평가 방법', '교육과정\n성취기준', '평가 시기']
    for r, label in enumerate(labels):
        header_cell(t.rows[r].cells[0], label)
    for i, ep in enumerate(eps):
        codes = '\n'.join(s.split(']')[0] + ']' for s in ep.get('standards', []))
        header_cell(t.rows[0].cells[i + 1], f'수행평가{i + 1}')
        fill_cell(t.rows[1].cells[i + 1], f'{round(100 / len(eps))}%', align='center')
        fill_cell(t.rows[2].cells[i + 1],
                  f"{ep.get('title', '')}\n{ep.get('method', '실기형')}",
                  align='center', valign='top')
        fill_cell(t.rows[3].cells[i + 1], codes, align='center', valign='top')
        fill_cell(t.rows[4].cells[i + 1], ep.get('period', eval_period), align='center')
    para(doc, '', after=6)

    # 성취율과 성취도
    para(doc, '3. 성취율과 성취도', size=10, bold=True, after=3)
    para(doc, '정기시험 및 수행평가의 반영비율 환산 점수의 합계(성취율)에 따라 다음과 같이 평정한다.', after=3)
    gw = [CW // 2, CW - CW // 2]
    t = make_table(doc, gw, 4)
    header_cell(t.rows[0].cells[0], '성취율')
    header_cell(t.rows[0].cells[1], '성취도')
    for r, (율, 도) in enumerate([('80%이상 ~ 100%', 'A'),
                                  ('60%이상 ~ 80%미만', 'B'),
                                  ('60%미만', 'C')], start=1):
        fill_cell(t.rows[r].cells[0], 율, align='center')
        fill_cell(t.rows[r].cells[1], 도, align='center', bold=True)
    para(doc, '', after=6)

    # 수행평가 세부 계획
    para(doc, '4. 수행평가 세부 계획', size=10, bold=True, after=4)
    order = ['가', '나', '다', '라', '마']
    sw = [1800, CW - 1800]
    for idx, ep in enumerate(eps):
        para(doc, f"{order[idx]}. {ep.get('title', '')}", size=9.5, bold=True,
             after=3, indent=280)
        stds = ep.get('standards', [''])
        elements = ep.get('elements', [])
        rows = len(stds) + 3  # 성취기준 n행 + 과제흐름 + 평가요소 + AI활용
        t = make_table(doc, sw, rows)
        # 성취기준 (세로 병합)
        header_cell(t.rows[0].cells[0], '성취기준')
        if len(stds) > 1:
            t.rows[0].cells[0].merge(t.rows[len(stds) - 1].cells[0])
        for i, s in enumerate(stds):
            fill_cell(t.rows[i].cells[1], s, valign='top')
        r = len(stds)
        header_cell(t.rows[r].cells[0], '수행 과제\n흐름(단계)')
        fill_cell(t.rows[r].cells[1], '\n'.join(ep.get('taskFlow', [])), valign='top')
        header_cell(t.rows[r + 1].cells[0], '평가요소\n(배점)')
        if elements:
            inner = t.rows[r + 1].cells[1].add_table(rows=1, cols=len(elements))
            inner.style = 'Table Grid'
            iw = (sw[1] - 200) // len(elements)
            for i, el in enumerate(elements):
                c = inner.rows[0].cells[i]
                c.width = Twips(iw)
                fill_cell(c, f"{el.get('name', '')}\n({el.get('score', '')}점)",
                          align='center')
        header_cell(t.rows[r + 2].cells[0], 'AI 활용')
        fill_cell(t.rows[r + 2].cells[1],
                  ep.get('aiPolicy', '본 수행평가는 실기 평가로 AI 도구 활용을 허용하지 않는다.'),
                  valign='top')
        para(doc, '', after=4)

        # 루브릭
        rubric = ep.get('rubric', [])
        total_rows = 1 + sum(len(r_.get('grades', [])) for r_ in rubric)
        rw = [2200, 900, CW - 3100]
        t = make_table(doc, rw, total_rows)
        header_cell(t.rows[0].cells[0], '평가요소(배점)')
        header_cell(t.rows[0].cells[1], '점수')
        header_cell(t.rows[0].cells[2], '채점 기준')
        row_i = 1
        for r_ in rubric:
            grades = r_.get('grades', [])
            start = row_i
            for g in grades:
                fill_cell(t.rows[row_i].cells[1], g.get('score', ''), align='center')
                fill_cell(t.rows[row_i].cells[2], g.get('criteria', ''), valign='top')
                row_i += 1
            fill_cell(t.rows[start].cells[0], r_.get('element', ''), align='center')
            if len(grades) > 1:
                t.rows[start].cells[0].merge(t.rows[row_i - 1].cells[0])
        para(doc, '[유의 사항] ※ 장기 결석, 백지 제출, 미완성 등 평가 결과 확인이 어려운 경우 별도 처리 기준에 따른다.',
             size=8, after=6)

    para(doc, '5. 유의 사항', size=10, bold=True, after=3)
    for line in 유의사항:
        para(doc, line, after=2)

    # ── 2부: 교수학습-평가 방법 (진도표) ──────────────
    page_break(doc)
    section_header(doc, '2', '(체육)과 교수학습-평가 방법')
    tw = [1600, 500, 2000, 1500, 1200, CW - 6800]
    rows2 = data.get('part2', {}).get('rows', [])
    t = make_table(doc, tw, len(rows2) + 1)
    for i, h in enumerate(['시기', '시수', '단원명(주제)\n[핵심 아이디어]', '성취기준',
                           '평가 유형', '평가와 연계한 수업 세부 방법']):
        header_cell(t.rows[0].cells[i], h)
    for r, row in enumerate(rows2, start=1):
        d = row.get('detail', {})
        det_parts = []
        for key, label in [('핵심개념', '핵심개념'), ('핵심탐구질문', '핵심(탐구)질문'),
                           ('개별화전략', '개별화 전략'), ('피드백전략', '피드백 전략'),
                           ('형성평가', '형성평가'), ('수행지시어', '수행지시어')]:
            if d.get(key):
                det_parts.append(f"[{label}]\n{d[key]}")
        fill_cell(t.rows[r].cells[0], row.get('period', ''), size=7.5, valign='top')
        fill_cell(t.rows[r].cells[1], row.get('hours', ''), align='center')
        fill_cell(t.rows[r].cells[2], row.get('unit', ''), size=7.5, valign='top')
        fill_cell(t.rows[r].cells[3], row.get('standard', ''), size=7.5, align='center')
        fill_cell(t.rows[r].cells[4], row.get('evalType', ''), size=7.5, align='center')
        fill_cell(t.rows[r].cells[5], '\n'.join(det_parts), size=7.5, valign='top')

    # ── 3부: 학기 단위 성취수준 ────────────────────────
    page_break(doc)
    section_header(doc, '3', '(체육)과 학기 단위 성취수준')
    para(doc, '1. 학기 단위 성취수준', size=10, bold=True, after=4)
    s3w = [700, 4000, CW - 4700]
    sem = data.get('part3', {}).get('semester', [])
    t = make_table(doc, s3w, len(sem) + 1)
    for i, h in enumerate(['성취수준', '성취수준 기술', '성취율']):
        header_cell(t.rows[0].cells[i], h)
    for r, s in enumerate(sem, start=1):
        header_cell(t.rows[r].cells[0], s.get('level', ''))
        fill_cell(t.rows[r].cells[1], s.get('desc', ''), valign='top')
        fill_cell(t.rows[r].cells[2], s.get('rate', ''), align='center')
    para(doc, '', after=6)

    para(doc, '2. 성취기준별 성취수준', size=10, bold=True, after=4)
    third = (CW - 3600) // 3
    s4w = [1600, 2000, third, third, CW - 3600 - 2 * third]
    by_std = data.get('part3', {}).get('byStandard', [])
    t = make_table(doc, s4w, len(by_std) + 1)
    for i, h in enumerate(['단원명', '교육과정 성취기준', 'A', 'B', 'C']):
        header_cell(t.rows[0].cells[i], h)
    for r, s in enumerate(by_std, start=1):
        fill_cell(t.rows[r].cells[0], s.get('unit', ''), size=7.5, align='center')
        fill_cell(t.rows[r].cells[1], s.get('standard', ''), size=7.5, valign='top')
        for i, lv in enumerate(['A', 'B', 'C']):
            fill_cell(t.rows[r].cells[2 + i], s.get(lv, ''), size=7.5, valign='top')

    doc.save(out_path)
    return out_path


def validate(data):
    """생성 전 일관성 검증. 문제 목록을 반환한다(비어 있으면 통과)."""
    problems = []
    eps = data.get('part1', {}).get('evalPlans', [])
    if not eps:
        problems.append('part1.evalPlans 가 비어 있음')
    for i, ep in enumerate(eps, 1):
        total = sum(el.get('score', 0) for el in ep.get('elements', []))
        rubric_els = {r.get('element') for r in ep.get('rubric', [])}
        plan_els = {el.get('name') for el in ep.get('elements', [])}
        if rubric_els != plan_els:
            problems.append(f'수행평가{i}: 평가요소와 루브릭 요소 불일치 {plan_els ^ rubric_els}')
        for r in ep.get('rubric', []):
            scores = [g.get('score', 0) for g in r.get('grades', [])]
            if scores != sorted(scores, reverse=True):
                problems.append(f"수행평가{i} 루브릭 '{r.get('element')}': 점수가 내림차순이 아님")
    rows = data.get('part2', {}).get('rows', [])
    meta_hours = data.get('meta', {}).get('totalHours')
    if meta_hours:
        s = sum(int(r.get('hours', 0)) for r in rows)
        if s != int(meta_hours):
            problems.append(f'진도표 시수 합계 {s} ≠ 총 차시 {meta_hours}')
    questions = [r.get('detail', {}).get('핵심탐구질문', '') for r in rows]
    dups = {q for q in questions if q and questions.count(q) > 1}
    if dups:
        problems.append(f'핵심탐구질문 중복: {len(dups)}건')
    sem = data.get('part3', {}).get('semester', [])
    if {s.get('level') for s in sem} != {'A', 'B', 'C'}:
        problems.append('part3.semester 에 A/B/C 3개 수준이 모두 필요')
    return problems


if __name__ == '__main__':
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)
    with open(sys.argv[1], encoding='utf-8') as f:
        raw = f.read()
    # 마크다운 코드블록·앞뒤 설명 제거 (Claude 답변 그대로 넣어도 동작)
    raw = raw.replace('```json', '').replace('```', '')
    raw = raw[raw.index('{'):raw.rindex('}') + 1]
    data = json.loads(raw)

    issues = validate(data)
    if issues:
        print('⚠ 검증 경고:')
        for p in issues:
            print('  -', p)

    sport = data.get('meta', {}).get('sport', '체육')
    out = sys.argv[2] if len(sys.argv) > 2 else f'체육과_평가운영계획_{sport}.docx'
    print('✓ 생성 완료:', build(data, out))
